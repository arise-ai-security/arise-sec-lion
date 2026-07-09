#!/usr/bin/env python
"""Generate the B1 subtree-modularization summary as a .docx.

Pure transform of the committed analysis artifacts — reads
``reports/modularity/aggregate.json``, converts the SVG figures to PNG with the
macOS-native ``qlmanage`` (no extra deps), and assembles a Word document with:
legend descriptions for each figure, plain-language definitions illustrated with
ONE worked run, and the most impactful data-backed findings (each tied to a
table and/or figure). No live DB needed.

    uv run --with python-docx python experiments/b1-batch-autogen/scripts/build_summary_docx.py
"""

from __future__ import annotations

import json
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

REPORT_DIR = Path(__file__).resolve().parent.parent / "reports" / "modularity"
FIG_DIR = REPORT_DIR / "figures"
OUT = REPORT_DIR / "B1_modularization_summary.docx"

# One representative run used to make every definition concrete (computed live;
# see commit history). exiv2.cve-2017-14857, run 62353082-… (20 nodes, 37 edges).
EX = "exiv2.cve-2017-14857"

FIGURES = [
    ("fig1_node_dsm.svg", "Role-resolved interaction DSM (all-runs average)",
     "Rows and columns are role slots, one per (module, tier): **bld·M / bld·W** = Builder "
     "manager-tier / worker-tier; **exp·M / exp·W** = Exploiter; **fix·M / fix·W** = Fixer; "
     "**rep·M** = Reporter (a single worker); **boss** = the Boss root. Each cell = the mean "
     "interaction weight per run between those two slots (messages + producer→consumer dataflow); "
     "darker = more interaction. Black lines separate modules.",
     "Dense dark blocks on the diagonal (manager↔worker inside one module) show cohesion; the "
     "near-empty off-diagonal shows low inter-module coupling; the boss row/column is the relay. A "
     "clean block-diagonal = modular."),
    ("fig2_module_dsm.svg", "Module×module interaction DSM (mean per run)",
     "Rows/cols are the five components builder / exploiter / fixer / reporter / boss. Each cell = "
     "mean interaction weight per run between the two modules; the diagonal is within-module "
     "(cohesion), off-diagonal is between-module (coupling), and the boss row/column is the relay.",
     "Bright diagonal + bright boss row/col, with a dim non-boss off-diagonal = cohesive modules "
     "whose only cross-talk is via the Boss."),
    ("fig3_module_dataflow_dsm.svg", "Directed producer→consumer dataflow (mean edges/run)",
     "A 4×4 directed matrix: the ROW is the producer (the module that wrote a file) and the COLUMN "
     "is the consumer (the module that later read it); each cell = mean such edges per run.",
     "Weight on the Builder→Exploiter→Fixer→Reporter direction (above the diagonal) means data flows "
     "along the prescribed pipeline; little back-flow."),
    ("fig4_recovery_nmi.svg", "Module recovery: detected communities vs the real modules",
     "Three bars = NMI (agreement, 0–1) between an unsupervised clustering of the interaction graph "
     "and the true module labels, for: the full graph, the dataflow-only (behavioral) graph, and a "
     "random-partition baseline. Whiskers = 95% CI of the mean.",
     "Full and dataflow-only bars far above the random baseline = the modules are recoverable from "
     "behavior alone; the gap between the first two = how much recovery is due to tree topology."),
    ("fig5_coupling_null_z.svg", "Inter-module coupling vs a permutation null (per-run z)",
     "Histogram of the per-run z-score of the observed cross-module dataflow fraction against a "
     "label-shuffling null. The dashed line is z = 0 (= chance).",
     "Mass to the LEFT of 0 = coupling lower than a random module assignment would give; further "
     "left = stronger (the per-run mean z is reported in Finding F2)."),
    ("fig6_module_instability.svg", "Per-module instability I = Ce/(Ca+Ce)",
     "One bar per module = mean instability: 0 = a pure producer others depend on (stable), 1 = a "
     "pure consumer that depends on others (unstable).",
     "The ascending Builder→Reporter staircase = the data-dependency direction mirrors the prescribed "
     "pipeline (Builder stable, Reporter unstable) — Conway's law."),
    ("fig7_cohesion_coupling.svg", "Per-module cohesion vs coupling (mean weight/run)",
     "Two bars per module: intra-module interaction (cohesion) vs inter-module interaction (coupling).",
     "Cohesion ≥ coupling for the three composite modules (Builder/Exploiter/Fixer). Reporter is a single "
     "node, so it has no internal cohesion and shows only coupling (a pure consumer)."),
    ("fig8_file_overlap_zone.svg", "Cross-module file interaction by workspace zone (mean/run)",
     "Per zone (`/src` target source, `/testcase` deliverables/handoff, `/work` scratch): write-write "
     "overlap, read-overlap, and producer→consumer edges, per run.",
     "Low write-write everywhere = little contention; read-overlap on `/src` = shared input (expected); "
     "producer→consumer on `/testcase`+`/src` = the prescribed handoff."),
]


def svg_to_png(svg: Path, out_dir: Path) -> Path:
    if not svg.exists():
        raise FileNotFoundError(f"missing figure: {svg}")
    subprocess.run(["qlmanage", "-t", "-s", "1600", "-o", str(out_dir), str(svg)],
                   check=True, capture_output=True)
    png = out_dir / (svg.name + ".png")
    if not png.exists():
        raise RuntimeError(f"qlmanage produced no PNG for {svg}")
    return png


def fmt(stat: dict) -> str:
    """median (IQR; mean, 95% CI). Median and mean are shown separately because the
    distributions are skewed, so the mean-CI need not bracket the median."""
    if not stat or stat.get("n", 0) == 0:
        return "n/a"
    ci = stat.get("ci95_mean") or [None, None]
    ci_txt = (f", 95% CI [{ci[0]:.3f}, {ci[1]:.3f}]") if ci[0] is not None else ""
    return (f"{stat['median']:.3f} (IQR {stat['iqr'][0]:.3f}–{stat['iqr'][1]:.3f}; "
            f"mean {stat['mean']:.3f}{ci_txt})")


def add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def finding(doc: Document, title: str, body: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor(0x0B, 0x53, 0x94)
    doc.add_paragraph(body)


def main() -> None:
    agg = json.loads((REPORT_DIR / "aggregate.json").read_text())
    c, perm = agg["coupling"], agg["permutation"]["dataflow_inter_fraction"]
    prof, flow = agg["module_profiles"], agg["module_dataflow_mean_per_run"]
    gwm, fo = agg["global_context"]["writes_by_module_kind"], agg["file_overlap"]
    art_total = sum(v.get("artifact_total", 0) for v in gwm.values())
    n = agg["n_runs"]
    mods = ["builder", "exploiter", "fixer", "reporter"]
    tmp = Path(tempfile.mkdtemp())

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading("Subtree Modularization of an Agentic Security-Analysis System", 0)
    doc.add_paragraph(
        f"Empirical analysis of the b1-batch-autogen study ({n} successful runs, one per CVE), "
        "grounded entirely in the event store. Each run decomposes a CVE-reproduction task into a tree: "
        "a Boss root spawns four phase modules — Builder, Exploiter, Fixer, Reporter — each a subtree of "
        "worker agents. We reconstruct each run's interaction graph from the events and measure, with "
        "software-architecture metrics, how modular the execution is.")

    doc.add_heading("1. Data & scope", 1)
    add_table(doc, ["stage", "runs", "note"], [
        ["B1 manifests", "124", "all runs tagged study_id=b1-batch-autogen"],
        ["unique CVEs (de-duplicated)", "112", "one run per CVE instance; 12 repeat runs dropped"],
        ["analyzed (success, complete tree)", str(n), "primary stratum; all 4 modules present"],
    ])
    doc.add_paragraph(
        "Source of truth is the Postgres event store; membership, run-id triangulation, and one-run-per-"
        "CVE de-duplication are enforced mechanically. Extraction of file touches was validated against "
        "on-disk diffs (Fixer patch-file recall ≈ 0.99).")

    doc.add_heading("2. Definitions (with a worked example)", 1)
    doc.add_paragraph(
        f"These are operationalizations computed from agent events — inspired by, but not identical to, "
        f"the textbook software-engineering metrics. Each is illustrated with one representative run, "
        f"{EX} (20 nodes, 37 interaction edges, Q = 0.44, recovery NMI = 0.89).")
    defs = [
        ("Node / Module",
         "A node is one agent (an event aggregate). Its module is the phase it belongs to, found by "
         "walking up to its depth-1 ancestor (the Boss's direct child) and reading that ancestor's "
         "[Builder]/[Exploiter]/[Fixer]/[Reporter] task prefix.",
         f"In {EX}: 20 nodes — a Boss, four module managers, and 15 workers split across the modules."),
        ("Interaction graph",
         "One vertex per node; an undirected edge weighted by how much two nodes interact = (#messages "
         "between them: parent→child briefings + child→parent reports) + (#producer→consumer file "
         "dependencies between them). Every coupling/modularity number is computed on this graph; the "
         "DSM figures draw it as a matrix.",
         f"In {EX}: 37 weighted edges over the 20 nodes."),
        ("Dataflow edge (producer→consumer)",
         "The behavioral dependency: if node A writes a file and node B later reads that same file "
         "(B≠A), record a directed edge A→B (\"B used what A produced\").",
         f"In {EX}: 5 cross-module edges, all Exploiter→Fixer — the Exploiter wrote "
         "/src/exiv2/src/image.cpp and an analysis note, and the Fixer read them."),
        ("Cohesion vs coupling",
         "Cohesion = total interaction-edge weight inside a module (its nodes talking to each other). "
         "Coupling = total edge weight between two different modules. Coupling ratio = "
         "coupling/(cohesion+coupling): the fraction of a module's interaction that crosses its boundary.",
         f"In {EX}: the Exploiter has cohesion 18 vs coupling 5 (ratio 0.22); the Fixer 20 vs 5 (0.20) — "
         "each talks to itself far more than to others."),
        ("Modularity Q + recovery (NMI/ARI)",
         "Q (Newman) scores how much edge weight falls inside the prescribed modules versus a random "
         "graph of the same degrees (>0 = clustered; ~0.3–0.7 = modular). Recovery asks the harder "
         "question: if we did NOT know the modules, would unsupervised clustering of the graph rediscover "
         "them? NMI and ARI (0–1, name-invariant) measure that agreement; a random-relabel baseline is "
         "the chance floor.",
         f"In {EX}: Q = 0.44, and clustering recovers the four modules at NMI = 0.89 / ARI = 0.88."),
        ("Instability I",
         "Per module, using only cross-module dataflow: Ca = #edges where the module is the WRITER "
         "(others depend on it ⇒ a stable producer); Ce = #edges where it is the READER (it depends on "
         "others ⇒ unstable). I = Ce/(Ca+Ce): 0 = pure producer, 1 = pure consumer.",
         f"In {EX}: the Exploiter is a pure producer (Ca=5, Ce=0, I=0.0) and the Fixer a pure consumer "
         "(Ca=0, Ce=5, I=1.0) — exactly the prescribed Exploiter→Fixer dependency."),
        ("Permutation null",
         "To test whether low coupling is below chance, we shuffle which module each node belongs to "
         "(keeping module sizes) 1000× and recompute the cross-module dataflow fraction; z = (observed − "
         "shuffled mean)/sd, with z≪0 meaning more self-contained than random.",
         f"In {EX}: 5 of 19 dataflow edges cross modules (fraction 0.26 — the statistic the null "
         "operates on), well under the shuffled-label mean."),
        ("Global-context write",
         "An entry a worker publishes to the shared blackboard: a DecisionRecorded (key→value+rationale) "
         "or ArtifactStored (key→description), attributed to the author's module. Measures reliance on "
         "shared state.",
         f"In {EX}: 78 writes — Exploiter 38, Fixer 26, Builder 11, Reporter 3."),
    ]
    for name, definition, example in defs:
        p = doc.add_paragraph()
        p.add_run(name + " — ").bold = True
        p.add_run(definition)
        ex = doc.add_paragraph(style="List Bullet")
        ex.add_run("Example: ").italic = True
        ex.add_run(example)

    doc.add_heading("3. Figures", 1)
    for svg_name, title, legend, how in FIGURES:
        doc.add_heading(title, 2)
        doc.add_picture(str(svg_to_png(FIG_DIR / svg_name, tmp)), width=Inches(6.0))
        lg = doc.add_paragraph()
        lg.add_run("Legend. ").bold = True
        lg.add_run(legend)
        rd = doc.add_paragraph()
        rd.add_run("How to read. ").bold = True
        rd.add_run(how)

    doc.add_heading("4. Findings (data-backed)", 1)

    finding(doc, "F1. The prescribed modules are recoverable from behavior.",
            "Unsupervised clustering of the interaction graph rediscovers the four modules far above "
            "chance — and still does on the purely-behavioral (dataflow-only) graph.")
    add_table(doc, ["recovery signal", "median (IQR; mean, 95% CI)"], [
        ["full interaction graph", fmt(c["recovery_nmi"])],
        ["dataflow-only (behavioral)", fmt(c["recovery_nmi_dataflow_only"])],
        ["random-partition baseline", fmt(c["recovery_nmi_random_baseline"])],
        ["ARI (full graph)", fmt(c["recovery_ari"])],
        ["modularity Q (labeled)", fmt(c["modularity_labeled"])],
    ])
    doc.add_paragraph("See Figure 4 (recovery) and Figure 1/2 (DSMs).")

    finding(doc, "F2. Inter-module coupling is below chance.",
            "The fraction of file-dependency edges that cross module boundaries is well below what a "
            "random module assignment produces (permutation null), on average.")
    add_table(doc, ["quantity", "value"], [
        ["observed cross-module dataflow fraction (median)", f"{perm['median_observed']:.3f}"],
        ["permutation-null fraction (median)", f"{perm['median_null_mean']:.3f}"],
        ["mean z-score", f"{perm['mean_z']:.2f}"],
        ["runs individually significant (p<0.05)", f"{perm['frac_p_below_0.05']*100:.0f}%"],
        ["inter-module interaction fraction (msgs+dataflow)", fmt(c["inter_module_fraction"])],
    ])
    doc.add_paragraph("See Figure 5 (per-run z histogram).")

    finding(doc, "F3. Each composite module is more cohesive than coupled; "
                 "F4. instability mirrors the prescribed pipeline (Conway's law).",
            "Within-module interaction dominates cross-module interaction for the three composite modules "
            "(Builder, Exploiter, Fixer); the Reporter is a single-node consumer with no internal cohesion "
            "(cohesion 0, coupling 1.0). Producer/consumer instability rises monotonically Builder→Reporter "
            "(0.00 → 0.04 → 0.88 → 1.00) — the data-dependency direction the Boss prescribed.")
    add_table(doc, ["module", "cohesion (intra)", "coupling (inter)", "coupling ratio", "instability I"],
              [[m, f"{prof[m]['intra_weight']:.1f}", f"{prof[m]['inter_weight']:.1f}",
                f"{prof[m]['coupling_ratio']:.2f}", f"{prof[m].get('instability', 0):.2f}"] for m in mods])
    doc.add_paragraph("See Figure 7 (cohesion vs coupling) and Figure 6 (instability).")

    finding(doc, "F5. Cross-module data flows along the prescribed handoff.",
            "Producer→consumer file dependencies concentrate on Exploiter→Fixer, matching the "
            "Builder→Exploiter→Fixer→Reporter DAG; inter-module *messages* are 0 by construction.")
    add_table(doc, ["producer → consumer", "mean edges/run"],
              [[k, f"{v}"] for k, v in flow.items() if k.split("->")[0] != k.split("->")[1]][:6])
    doc.add_paragraph(
        f"Inter-module messages across all {n} runs: {c['inter_module_message_count_total']} "
        "(a structural zero — briefings/reports only run along tree edges). See Figure 3.")

    finding(doc, "F6. The shared 'global context' is heavily used and Exploiter-dominated — "
                 "an annotation layer, not a payload store.",
            f"Used in {agg['global_context']['runs_with_any_writes']}/{n} runs "
            f"(median {agg['global_context']['writes_total']['median']:.0f} writes/run). Across all runs, "
            f"0 of {art_total} artifacts contain actual file content (they are prose pointers/summaries of "
            "workspace files), and ~75% of keys are run-unique/ad-hoc — so it functions as a "
            "human-readable changelog, not the machine-keyable table the prompts specify.")
    add_table(doc, ["module", "decisions/run", "artifacts/run", "total/run"],
              [[m, f"{gwm.get(m, {}).get('decision_per_run', 0)}",
                f"{gwm.get(m, {}).get('artifact_per_run', 0)}",
                f"{round(gwm.get(m, {}).get('decision_per_run', 0) + gwm.get(m, {}).get('artifact_per_run', 0), 2)}"]
               for m in mods])

    finding(doc, "F7. Modules share recon reads (no strict task de-duplication); "
                 "F8. file contention is minimal and matches the handoff.",
            "Modules re-read overlapping source files during analysis (so 'no duplicated analysis' does "
            "not hold strictly — but the overlap is shared input on /src). Files written by ≥2 modules "
            "are few and are almost all the Exploiter+Fixer pair co-editing the vulnerable source.")
    add_table(doc, ["metric", "value"], [
        ["recon-read Jaccard between modules", fmt(agg["task_overlap"]["recon_read_mean_jaccard"])],
        [f"write-write overlaps by zone (total over {n} runs)", str(fo["write_write_overlap_by_zone_total"])],
        ["write-write by module pair", str(fo["write_write_pairs_total"])],
        ["producer→consumer edges by zone (total)", str(fo["producer_consumer_edges_by_zone_total"])],
        ["read overlap by zone (total)", str(fo["read_overlap_by_zone_total"])],
    ])
    doc.add_paragraph("See Figure 8 (file interaction by zone).")

    doc.add_heading("5. Reading of the evidence & limitations", 1)
    bullet(doc, "Supported: the system executes a clean, recoverable, low-coupling modular decomposition "
                "whose data-dependency direction mirrors the prescribed pipeline.")
    bullet(doc, "Not supported as stated: 'no duplicated analysis' (modules share recon reads) and "
                "'minimal global-context reliance' (it is heavily used).")
    bullet(doc, "The module partition is prescribed by the Boss prompt, so this is execution fidelity + "
                "recoverability, not spontaneous emergence.")
    bullet(doc, "Inter-module messaging is 0 by construction (the schema cannot encode it); the "
                f"interaction graph is message-dominated, so dataflow-only recovery "
                f"({c['recovery_nmi_dataflow_only']['median']:.2f} NMI) is the behavioral lower bound. "
                "Shell-mediated writes are under-captured (low confidence).")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
